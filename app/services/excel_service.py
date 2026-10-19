
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.user import User
from app.models.crop import Crop
from app.models.order import Order
from app.utils.ist_time import format_ist_datetime


def apply_table_styling(ws, header_color="2D6A4F"):
    """Format the spreadsheet for easy reading."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        fill_type="solid",
        fgColor=header_color,
    )
    thin_border = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border

    ws.freeze_panes = "A2"

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    for column in ws.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )
        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(
            max(max_length + 3, 12),
            45,
        )


def _create_excel(headers, rows, sheet_name, header_color):
    """Create an Excel file in memory and return its byte stream."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    worksheet.append(headers)

    for row in rows:
        worksheet.append(row)

    apply_table_styling(worksheet, header_color)

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    return stream


class ExcelService:

    @staticmethod
    def export_users(db: Session) -> io.BytesIO:
        """Export registered users without password information."""
        users = db.query(User).order_by(User.id).all()

        headers = [
            "User ID",
            "Full Name",
            "Role",
            "Email",
            "Phone",
            "Address",
            "Account Status",
            "Registration Date (IST)",
        ]

        rows = [
            [
                user.id,
                user.full_name,
                user.role,
                user.email,
                user.phone or "",
                user.address or "",
                user.status,
                format_ist_datetime(user.created_at),
            ]
            for user in users
        ]

        return _create_excel(
            headers, rows, "Registered Users", "1E40AF"
        )

    @staticmethod
    def export_crops(db: Session) -> io.BytesIO:
        """Export crop listings with associated farmer information."""
        crops = db.query(Crop).order_by(Crop.crop_id).all()

        headers = [
            "Crop ID",
            "Farmer ID",
            "Farmer Name",
            "Crop Name",
            "Category",
            "Description",
            "Available Quantity",
            "Unit",
            "Price Per Unit (INR)",
            "AI Suggested Price",
            "Quality Grade",
            "Freshness Score",
            "Location",
            "Listing Date (IST)",
            "Availability",
        ]

        rows = []

        for crop in crops:
            farmer = crop.farmer

            rows.append([
                crop.crop_id,
                crop.farmer_id,
                farmer.full_name if farmer else "Unknown",
                crop.crop_name,
                crop.category,
                crop.description or "",
                crop.quantity_available,
                crop.unit,
                crop.farmer_price,
                crop.ai_suggested_price,
                crop.quality_grade,
                crop.freshness_score,
                crop.location or "",
                format_ist_datetime(crop.created_at),
                crop.status,
            ])

        return _create_excel(
            headers, rows, "Farmer Crops", "2D6A4F"
        )

    @staticmethod
    def export_orders(db: Session) -> io.BytesIO:
        """Export order records using the existing order model."""
        orders = (
            db.query(Order)
            .order_by(Order.order_id.desc())
            .all()
        )

        headers = [
            "Order ID",
            "Order Date & Time (IST)",
            "Consumer Name",
            "Consumer Phone",
            "Farmer Name",
            "Crop Name",
            "Quantity",
            "Unit",
            "Price Per Unit (INR)",
            "Subtotal (INR)",
            "Delivery Charge (INR)",
            "Discount (INR)",
            "Total Amount (INR)",
            "Payment Method",
            "Payment Status",
            "Delivery Partner",
            "Order Status",
            "Delivery Address",
        ]

        rows = []

        for order in orders:
            consumer = order.consumer
            farmer = order.farmer
            crop = order.crop
            delivery_partner = order.delivery_partner

            rows.append([
                order.order_id,
                format_ist_datetime(order.order_date),
                consumer.full_name if consumer else "",
                consumer.phone if consumer else "",
                farmer.full_name if farmer else "",
                crop.crop_name if crop else "",
                order.quantity,
                crop.unit if crop else "",
                order.price_per_kg,
                order.subtotal,
                order.delivery_charge,
                order.discount,
                order.total_amount,
                order.payment_method,
                order.payment_status,
                (
                    delivery_partner.full_name
                    if delivery_partner else "Unassigned"
                ),
                order.status,
                order.delivery_address,
            ])

        return _create_excel(
            headers, rows, "Platform Orders", "7C3AED"
        )
